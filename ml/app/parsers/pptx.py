import io

from pptx import Presentation

from app.parsers.base import ParseError


def parse(data: bytes) -> str:
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"pptx: cannot read presentation: {e}") from e

    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return "\n\n".join(parts)